# -*- coding: utf-8 -*-
"""专家盲评分析（按 papers/rasfr/expert_review/protocol.md 的预定计划）。

  venv/bin/python -m score_learning.induction.expert_analyse

读入 约束盲评表_<评审人>.xlsx 与 runs/induction/paper/expert_key.json，
输出 runs/induction/paper/expert_results.json。
"""
from __future__ import annotations

import glob
import json
import os
import random
from collections import Counter
from typing import Dict, List

from openpyxl import load_workbook

KEY = 'runs/induction/paper/expert_key.json'
SHEETS = 'papers/rasfr/expert_review/约束盲评表_*.xlsx'
OUT = 'runs/induction/paper/expert_results.json'
CATS = ('有效', '合理', '风格依赖', '无效')
ACCEPT = ('有效', '合理')
METHODS = ('ras', 'pmi', 'rarity', 'productive')
N_BOOT = 2000


def load_ratings() -> Dict[str, Dict[str, str]]:
    out = {}
    for path in sorted(glob.glob(SHEETS)):
        rater = os.path.splitext(os.path.basename(path))[0].split('_')[-1]
        ws = load_workbook(path)['评分']
        out[rater] = {r[0].value: (r[4].value or '无法判断') for r in ws.iter_rows(min_row=2) if r[0].value}
    return out


def kripp_alpha(units: List[List[str]]) -> float:
    """名义尺度 Krippendorff alpha；units 为每个条目的有效评分列表。"""
    o = Counter()
    for vals in units:
        m = len(vals)
        if m < 2:
            continue
        for i, a in enumerate(vals):
            for j, b in enumerate(vals):
                if i != j:
                    o[(a, b)] += 1.0 / (m - 1)
    n_c = Counter()
    for (a, _b), w in o.items():
        n_c[a] += w
    n = sum(n_c.values())
    if n <= 1:
        return float('nan')
    d_o = sum(w for (a, b), w in o.items() if a != b) / n
    d_e = sum(n_c[a] * n_c[b] for a in n_c for b in n_c if a != b) / (n * (n - 1))
    return 1.0 - d_o / d_e if d_e > 0 else float('nan')


def cohen_kappa(x: List[str], y: List[str]) -> float:
    n = len(x)
    if n == 0:
        return float('nan')
    po = sum(a == b for a, b in zip(x, y)) / n
    cx, cy = Counter(x), Counter(y)
    pe = sum(cx[c] * cy[c] for c in set(cx) | set(cy)) / n ** 2
    return (po - pe) / (1 - pe) if pe < 1 else float('nan')


def main():
    key = json.load(open(KEY))['items']
    first = [it for it in key if it['repeat_of'] is None]
    repeats = [it for it in key if it['repeat_of']]
    ratings = load_ratings()
    raters = sorted(ratings)
    res = {'raters': raters, 'n_items': len(first), 'n_repeats': len(repeats)}

    nominal = [[ratings[r][it['item']] for r in raters if ratings[r][it['item']] in CATS] for it in first]
    binary = [['A' if v in ACCEPT else 'N' for v in u] for u in nominal]
    res['alpha_nominal'] = kripp_alpha(nominal)
    res['alpha_binary'] = kripp_alpha(binary)
    res['rating_dist'] = {r: dict(Counter(ratings[r].values())) for r in raters}

    intra = {}
    for r in raters:
        a = [ratings[r][it['repeat_of']] for it in repeats]
        b = [ratings[r][it['item']] for it in repeats]
        intra[r] = {'identical': sum(x == y for x, y in zip(a, b)) / len(a), 'kappa': cohen_kappa(a, b)}
    res['intra_rater'] = intra

    def score(it, rs=raters):
        v = [ratings[r][it['item']] for r in rs if ratings[r][it['item']] in CATS]
        return sum(x in ACCEPT for x in v) / len(v) if v else None

    items = []
    for it in first:
        s = score(it)
        items.append({'item': it['item'], 'polarity': it['polarity'], 'methods': it['methods'],
                      'reference': bool(it['matches_reference']), 'score': s,
                      'endorsed': s is not None and s > 0.5,
                      'class': _cls(it)})
    res['items'] = items

    def validity(its, rs):
        out = {}
        for m in METHODS:
            for pol in ('prohibition', 'preference', 'all'):
                sel = [i for i in its if m in i['methods'] and (pol == 'all' or i['polarity'] == pol)]
                sc = [score(next(k for k in first if k['item'] == i['item']), rs) for i in sel]
                sc = [x for x in sc if x is not None]
                out[f'{m}/{pol}'] = sum(sc) / len(sc) if sc else None
        return out

    res['validity'] = validity(items, raters)
    rng = random.Random(2026)
    boot = []
    for _ in range(N_BOOT):
        rs = [rng.choice(raters) for _ in raters]
        its = [rng.choice(items) for _ in items]
        boot.append(validity(its, rs))
    ci = {}
    for m in METHODS[1:]:
        for pol in ('prohibition', 'preference', 'all'):
            d = [b[f'ras/{pol}'] - b[f'{m}/{pol}'] for b in boot
                 if b[f'ras/{pol}'] is not None and b[f'{m}/{pol}'] is not None]
            d.sort()
            ci[f'ras-{m}/{pol}'] = {'diff': res['validity'][f'ras/{pol}'] - res['validity'][f'{m}/{pol}'],
                                    'ci': (d[int(0.025 * (len(d) - 1))], d[int(0.975 * (len(d) - 1))])}
    res['validity_diff'] = ci

    tab = Counter((i['reference'], i['endorsed']) for i in items if i['score'] is not None)
    res['oracle_crosstab'] = {'ref_endorsed': tab[(True, True)], 'ref_not_endorsed': tab[(True, False)],
                              'noref_endorsed': tab[(False, True)], 'noref_not_endorsed': tab[(False, False)]}
    sel = [i for i in items if i['score'] is not None]
    res['oracle_kappa'] = cohen_kappa([str(i['reference']) for i in sel], [str(i['endorsed']) for i in sel])
    per_m = {}
    for m in METHODS:
        sel = [i for i in items if m in i['methods'] and i['score'] is not None]
        per_m[m] = {'n': len(sel), 'endorsed': sum(i['endorsed'] for i in sel),
                    'reference': sum(i['reference'] for i in sel),
                    'endorsed_not_reference': sum(i['endorsed'] and not i['reference'] for i in sel)}
    res['per_method'] = per_m
    by_cls = {}
    for c in sorted({i['class'] for i in items}):
        sc = [i['score'] for i in items if i['class'] == c and i['score'] is not None]
        by_cls[c] = {'n': len(sc), 'mean_score': sum(sc) / len(sc) if sc else None}
    res['by_provenance'] = by_cls

    with open(OUT, 'w') as fh:
        json.dump(res, fh, indent=1, ensure_ascii=False)
    print(json.dumps({k: res[k] for k in ('raters', 'alpha_nominal', 'alpha_binary', 'intra_rater', 'validity',
                                          'validity_diff', 'oracle_crosstab', 'oracle_kappa', 'per_method',
                                          'by_provenance')}, indent=1, ensure_ascii=False, default=str))


WITHIN = {'M', 'G', 'D', 'P', 'Q'}


def _cls(it) -> str:
    fld = it['cells'][0].split(':', 1)[0]
    if fld in WITHIN:
        return 'within-voice/positional'
    zr = max(it['signed_z_real'].values())
    zd = max(it['signed_z_desync'].values())
    if zr <= 0:
        return 'no replay evidence'
    return 'relational' if zd < 0.5 * zr else 'persists without coordination'


if __name__ == '__main__':
    main()
