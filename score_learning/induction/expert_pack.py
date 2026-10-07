# -*- coding: utf-8 -*-
"""专家盲评材料包：四种方法在留出 A 上各取前 N 条禁止/偏好，合并去重、打乱、去方法标签。

  venv/bin/python -m score_learning.induction.expert_pack

给专家：papers/rasfr/expert_review/约束盲评表.xlsx（说明、评分、字段说明三页）。
不给专家：runs/induction/paper/expert_key.json（条目 → 方法与名次、参考约束匹配、溯源 z）。
"""
from __future__ import annotations

import json
import os
import random
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from score_learning.induction import bootstrap as bs
from score_learning.induction import fri
from score_learning.induction.baselines import EVAL_FIELDS
from score_learning.induction.fresh_eval import METHODS
from score_learning.induction.fri import MIN_NULL, SPARSE_FIELDS, _ranked, annotate_z
from score_learning.induction.paper import _desync_piece, _git_head, load_chorale_pool

TOP_N = 20
N_REPEAT = 10
SEED = 2026
XLSX = 'papers/rasfr/expert_review/约束盲评表.xlsx'
KEY = 'runs/induction/paper/expert_key.json'

DEG = {'0': '1', '1': '♭2', '2': '2', '3': '♭3', '4': '3', '5': '4', '6': '♯4', '7': '5',
       '8': '♭6', '9': '6', '10': '♭7', '11': '7', 'x': '其他（变化音）', 'none': '无'}
VOICE = {'S': '女高', 'A': '女中', 'T': '男高', 'B': '男低'}
ROLE = {'outer': '外声部（女高—男低）', 'adjacent': '相邻声部', 'diagonal': '隔一声部'}
MOTION = {'same': '同向', 'opp': '反向', 'obl': '斜向（一声部不动）', 'stat': '两声部都不动'}
GAP = {'cross': '交叉（上方声部低于下方声部）', 'le8': '0–7 半音', '8to12': '8–12 半音',
       '12to19': '13–19 半音', 'gt19': '超过 19 半音'}
MOVE = {'+1': '上行小二度', '-1': '下行小二度', '-2': '下行大二度', 'stay': '不动', 'step': '上行大二度',
        'tritone': '三全音', '7th': '七度（10–11 半音）', 'oct+': '八度或更大', 'leap': '其他跳进（3–9 半音，三全音除外）'}
LEAP = {'small': '级进或小跳（1–4 半音）', 'leap': '跳进（5–9 半音，三全音除外）', 'tritone': '三全音',
        '7th': '七度（10–11 半音）', 'oct+': '八度或更大'}
D_BITS = [('st01', '进入中音为级进'), ('st12', '离开中音为级进'), ('same', '进入与离开同方向'),
          ('ret', '回到前一个音'), ('inn', '中音与其他声部同时音同音级（和弦音）'),
          ('held', '中音被延住而其他声部换音'), ('leap', '跳进（≥3 半音）进入中音'),
          ('ant', '中音不属于当前和音而属于下一和音（先现）'),
          ('holdO', '中音出现时其他声部全部保持'), ('hov', '跨拍延音形态')]
YES = {'1': '是', '0': '否'}

FIELD_DOC = [
    ('声部对进行', '相邻两个完整切片（四声部同时发声）之间，一对声部的运动：是否都动、移动量是否相同、新音程是否为纯一/八度或纯五度、音程是否保持、上方声部是否跳进、运动方向、声部对类型。'),
    ('声部超越', '相邻声部中，一个声部的新音是否越过另一声部的前一个音。'),
    ('相邻声部间距', '同一时刻相邻两声部之间的距离（半音数，分五档）。'),
    ('旋律进行', '单一声部换音时的音程，区分外声部/内声部、出发音是否为导音（7 级）或降 7 级。'),
    ('和音构成', '一个四声部和音中：导音出现次数、是否有声部与低音构成三度、降 7 级出现次数、是否包含主三和弦三个音。'),
    ('重复音级', '一个四声部和音中被重复的音级（若有多个，取级数最小者）。'),
    ('全曲结尾', '全曲最后两个完整切片的低音与女高音级。'),
    ('乐句结尾', '每个乐句最后两个不同切片的低音与女高音级；区分全曲最后一句与中间乐句。乐句以女高长音为界自动切分。'),
    ('和音连接', '相继两个不同音集的低音级、音集（以主音为 1 的音级集合）、是否含变化音，及是否位于乐句末尾（每句最后三个音集之内）；较粗的一种只记低音级进行。'),
    ('三音形态', '某一声部三个相继音（前—中—后）的形态：进入/离开是否级进、方向、是否回原音、中音是否为和弦音、延留、先现、持续音等。'),
]

RATINGS = ['有效：公认规范或教科书规则', '合理：可解释的风格倾向', '风格依赖：只在部分风格/时期成立',
           '无效：不成立、错误或无音乐意义', '无法判断：描述不清或信息不足']


def _deg(v: str) -> str:
    return DEG.get(v, v)


def _set(s: str) -> str:
    return '{' + '、'.join(_deg(d) for d in s.split('.')) + '}'


def _kv(cell: str) -> Dict[str, str]:
    out = {}
    for part in cell.split('|'):
        k, _, v = part.partition('=')
        out[k if v else 'pair'] = v if v else k
    return out


def _sonority(tok: str) -> str:
    d = _kv(tok)
    s = f'低音 {_deg(d["b"])}'
    if 'set' in d:
        s += f'，音集 {_set(d["set"])}' + ('，含变化音' if d.get('ch') == '1' else '')
    return s


def describe(fld: str, cell: str) -> str:
    if fld in ('K', 'J', 'Q'):
        a, b = cell.split('/to/')
        b, _, slot = b.rpartition('|slot=')
        where = {'K': {'fin': '位于乐句末尾（每句最后三个音集之内）', 'int': '位于乐句中间'},
                 'J': {'fin': '位于乐句末尾（每句最后三个音集之内）', 'int': '位于乐句中间'},
                 'Q': {'fin': '全曲最后一句', 'int': '中间乐句'}}[fld][slot]
        if fld == 'Q':
            da, db = _kv(a), _kv(b)
            return (f'乐句结尾最后两个不同切片：低音 {_deg(da["b"])}/女高 {_deg(da["sop"])} → '
                    f'低音 {_deg(db["b"])}/女高 {_deg(db["sop"])}；{where}')
        kind = '和音连接' if 'set=' in a else '低音级进行'
        return f'{kind}：{_sonority(a)} → {_sonority(b)}；{where}'
    d = _kv(cell)
    if fld == 'V':
        parts = [ROLE[d['role']], '两声部都动' if d['both'] == '1' else '至少一个声部不动',
                 '两声部移动量相同' if d['eqΔ'] == '1' else '两声部移动量不同',
                 '到达的音程为纯一/八度或纯五度（含复音程）' if d['pc07'] == '1' else '到达的音程不是纯一/八度或纯五度',
                 '音程保持不变' if d['hold'] == '1' else '音程改变',
                 '上方声部跳进（>2 半音）' if d['leapU'] == '1' else '上方声部级进或不动',
                 '运动方向：' + MOTION[d['sgn']]]
        s = '声部对从一个完整切片到下一个：' + '；'.join(parts)
        if d['both'] == '1' and d['eqΔ'] == '1' and d['pc07'] == '1':
            s += '（即平行纯五度或纯八度/同度）'
        elif d['both'] == '1' and d['pc07'] == '1' and d['sgn'] == 'same':
            s += '（即同向进入纯五度或纯八度）'
        return s
    if fld == 'O':
        a, b = d['pair'][0], d['pair'][1]
        return (f'相邻声部 {VOICE[a]}—{VOICE[b]}：一声部的新音'
                + ('越过另一声部的前一个音（声部超越）' if d['ov'] == '1' else '没有越过另一声部的前一个音'))
    if fld == 'H':
        a, b = d['pair'][0], d['pair'][1]
        return f'同一时刻 {VOICE[a]} 与 {VOICE[b]} 的距离：{GAP[d["gap"]]}'
    if fld == 'M':
        frm = '导音（7 级）' if d['lt'] == '1' else ('降 7 级' if d['b7'] == '1' else '非导音、非降 7 级的音')
        return (f'{"外声部（女高或男低）" if d["out"] == "1" else "内声部（女中或男高）"}从{frm}出发的旋律进行：'
                f'{MOVE[d["mv"]]}')
    if fld == 'G':
        return f'任一声部换音时的旋律音程：{LEAP[cell]}'
    if fld == 'S':
        return (f'一个四声部和音中：导音出现 {d["nlt"]} 次{"或以上" if d["nlt"] == "2" else ""}；'
                + ('没有任何声部与低音构成三度（或复三度）' if d['miss3'] == '1' else '有声部与低音构成三度') + '；'
                + f'降 7 级出现 {d["n7"]} 次{"或以上" if d["n7"] == "2" else ""}；'
                + ('不包含主三和弦的全部三个音' if d['missTon'] == '1' else '包含主三和弦的全部三个音'))
    if fld == 'U':
        return '一个四声部和音中' + ('没有重复的音级' if d['dbl'] == 'none'
                                    else f'被重复的音级（取级数最小者）为 {_deg(d["dbl"])}')
    if fld == 'P':
        return f'全曲最后两个完整切片之一：低音 {_deg(d["b"])}，女高 {_deg(d["sop"])}'
    if fld == 'D':
        return '某声部三个相继音（前—中—后）：' + '；'.join(f'{t}：{YES[d[k]]}' for k, t in D_BITS)
    raise ValueError(fld)


def top_items(fields, n=TOP_N, seed=SEED) -> Dict[str, List[tuple]]:
    """各极性按带符号 z 排序取前 n；并列（稀有度基线常见）按固定种子随机打破。"""
    z = annotate_z(fields)
    rng = random.Random(seed)
    out = {}
    for pol, sgn in (('prohibition', -1.0), ('preference', 1.0)):
        scored = []
        for fld in EVAL_FIELDS:
            for c in (r[0] for r in _ranked(fields.get(fld, []), 3 if fld in SPARSE_FIELDS else MIN_NULL)):
                scored.append((sgn * z[fld][c], rng.random(), fld, c))
        scored.sort(reverse=True)
        out[pol] = [(fld, c, s) for s, _t, fld, c in scored[:n]]
    return out


def write_xlsx(items: List[dict], path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = '说明'
    lines = [
        '四声部众赞歌约束盲评',
        '',
        '下表每一行是一条从巴赫四声部众赞歌中统计得到的陈述：某种情形“被回避”（出现得比预期少）或“被偏好”（出现得比预期多）。',
        '条目来自若干种不同的统计方法，已打乱顺序并去掉来源标记；部分条目会重复出现，用于检验评分一致性。',
        '请仅根据您对该风格（巴赫四声部众赞歌 / 共性写作时期四部和声）的专业判断评分，不需要查阅本研究的任何材料。',
        '',
        '评分（“评分”页 E 列，下拉选择）：',
        *[f'  {r}' for r in RATINGS],
        '',
        '“有效”与“合理”的区别：有效指和声或对位教材中通常作为规则或明确规范陈述的内容；合理指教材不一定写成规则、但您认为符合该风格写作习惯的倾向。',
        '若一条陈述的方向（回避/偏好）与您的判断相反，请评为“无效”，并在备注中说明。',
        '音级以主音为 1 计（大调与小调同）：3 为大三度音，♭3 为小三度音，7 为导音，♭7 为降七级；“其他（变化音）”指七个自然音级以外的音。',
        '“完整切片”指四个声部同时发声的时刻；每一次任一声部换音即开始一个新切片。',
        '各字段的含义见“字段说明”页。',
        '',
        '请独立评分，评分期间不要与其他评审者讨论。预计用时 60–90 分钟。',
    ]
    for i, t in enumerate(lines, 1):
        ws.cell(row=i, column=1, value=t).alignment = Alignment(wrap_text=True, vertical='top')
    ws['A1'].font = Font(bold=True, size=14)
    ws.column_dimensions['A'].width = 110

    ws = wb.create_sheet('评分')
    head = ['编号', '方向', '类别', '陈述', '评分', '备注']
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill('solid', fgColor='DDDDDD')
    for it in items:
        ws.append([it['item'], it['direction'], it['category'], it['statement'], None, None])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    dv = DataValidation(type='list', formula1='"' + ','.join(r.split('：')[0] for r in RATINGS) + '"',
                        allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'E2:E{len(items) + 1}')
    for col, w in zip('ABCDEF', (8, 8, 12, 80, 12, 30)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'

    ws = wb.create_sheet('字段说明')
    ws.append(['类别', '含义'])
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in FIELD_DOC:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 100
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)


CATEGORY = {'V': '声部对进行', 'O': '声部超越', 'H': '相邻声部间距', 'M': '旋律进行', 'G': '旋律进行',
            'S': '和音构成', 'U': '重复音级', 'P': '全曲结尾', 'Q': '乐句结尾', 'K': '和音连接',
            'J': '和音连接', 'D': '三音形态'}


def main():
    with open(bs.FINAL) as fh:
        fin = json.load(fh)
    by, _usable, _ = load_chorale_pool()
    tr = [by[i] for i in fin['train_ids']]
    hold = [by[i] for i in fin['hold_ids']]
    rng_d = random.Random(99)
    hold_d = [_desync_piece(p, rng_d) for p in hold]
    f, fd = fri.holdout_fields(tr, hold, seed=0), fri.holdout_fields(tr, hold_d, seed=0)
    zr, zd = annotate_z(f), annotate_z(fd)
    preds = fri.book_predicates()

    merged: Dict[tuple, dict] = {}
    for m, t in METHODS.items():
        for pol, rows in top_items(t(f)).items():
            for rank, (fld, cell, _s) in enumerate(rows, 1):
                stmt = describe(fld, cell)
                e = merged.setdefault((pol, stmt), {'polarity': pol, 'statement': stmt, 'category': CATEGORY[fld],
                                                    'cells': set(), 'ranks': {}})
                e['cells'].add((fld, cell))
                e['ranks'][m] = min(rank, e['ranks'].get(m, rank))

    rng = random.Random(SEED)
    items = list(merged.values())
    rng.shuffle(items)
    repeats = rng.sample(range(len(items)), N_REPEAT)
    order = items + [items[i] for i in repeats]
    tail = order[len(items):]
    rng.shuffle(tail)
    order = order[:len(items)] + tail
    sheet, key = [], []
    for i, it in enumerate(order, 1):
        first = items.index(it) + 1
        sheet.append({'item': f'C{i:03d}', 'direction': '回避' if it['polarity'] == 'prohibition' else '偏好',
                      'category': it['category'], 'statement': it['statement']})
        sgn = -1.0 if it['polarity'] == 'prohibition' else 1.0
        key.append({
            'item': f'C{i:03d}', 'repeat_of': None if i == first else f'C{first:03d}',
            'polarity': it['polarity'], 'statement': it['statement'],
            'cells': sorted(f'{a}:{b}' for a, b in it['cells']),
            'methods': it['ranks'],
            'matches_reference': sorted({n for fld, cell in it['cells'] for n, fp, p, pred in preds
                                         if fp == fld and p == ('under' if sgn < 0 else 'over')
                                         and pred(cell, fld)}),
            'signed_z_real': {f'{a}:{b}': round(sgn * zr[a][b], 2) for a, b in it['cells']},
            'signed_z_desync': {f'{a}:{b}': round(sgn * zd.get(a, {}).get(b, 0.0), 2) for a, b in it['cells']},
        })
    write_xlsx(sheet, XLSX)
    with open(KEY, 'w') as fh:
        json.dump({'commit': _git_head(), 'corpus': 'hold-out A', 'top_n': TOP_N, 'seed': SEED,
                   'n_unique': len(items), 'n_repeats': N_REPEAT, 'items': key}, fh, indent=1, ensure_ascii=False)
    per = {m: sum(m in it['ranks'] for it in items) for m in METHODS}
    shared = sum(len(it['ranks']) > 1 for it in items)
    print(f'{len(items)} unique items (+{N_REPEAT} repeats); per method {per}; shared by >1 method {shared}')
    print('wrote', XLSX, 'and', KEY)


if __name__ == '__main__':
    main()
